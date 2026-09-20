"""Safe, typed attachment readers for the four MVP formats."""
from __future__ import annotations

from pathlib import Path

from .models import DocumentReadResult, TextBlock


def _failure(path: Path, parser: str, status: str, code: str) -> DocumentReadResult:
    return DocumentReadResult(
        path=str(path),
        filename=path.name,
        readability_status=status,
        parser_name=parser,
        parser_confidence=0.0,
        error_code=code,
    )


def _finalize(path: Path, parser: str, confidence: float, blocks, tables=None, metadata=None):
    blocks = [block for block in blocks if block.text.strip()]
    if not blocks:
        return _failure(path, parser, "UNREADABLE", "NO_READABLE_TEXT")
    return DocumentReadResult(
        path=str(path),
        filename=path.name,
        text_blocks=blocks,
        tables=tables or [],
        metadata={"size_bytes": path.stat().st_size, **(metadata or {})},
        readability_status="READABLE",
        parser_name=parser,
        parser_confidence=confidence,
    )


def _read_txt(path: Path) -> DocumentReadResult:
    text = path.read_text(encoding="utf-8", errors="replace")
    return _finalize(path, "txt", 0.99, [TextBlock(text=text, location={"line_start": 1})])


def _read_pdf(path: Path) -> DocumentReadResult:
    import logging

    from pypdf import PdfReader

    logging.getLogger("pypdf").setLevel(logging.ERROR)
    reader = PdfReader(str(path))
    blocks = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        blocks.append(TextBlock(text=text, location={"page": page_number}))
    return _finalize(path, "pypdf", 0.88, blocks, metadata={"page_count": len(reader.pages)})


def _read_docx(path: Path) -> DocumentReadResult:
    from docx import Document

    document = Document(str(path))
    blocks = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        if paragraph.text.strip():
            blocks.append(TextBlock(paragraph.text, {"paragraph": index}))

    tables: list[list[list[str]]] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row_index, row in enumerate(table.rows, start=1):
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                rows.append(values)
                blocks.append(TextBlock(" | ".join(values), {"table": table_index, "row": row_index}))
        if rows:
            tables.append(rows)
    return _finalize(
        path,
        "python-docx",
        0.90,
        blocks,
        tables=tables,
        metadata={"paragraph_count": len(document.paragraphs), "table_count": len(document.tables)},
    )


def _read_xlsx(path: Path) -> DocumentReadResult:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks = []
    tables: list[list[list[str]]] = []
    for sheet in workbook.worksheets:
        rows = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = ["" if value is None else str(value).strip() for value in row]
            while values and not values[-1]:
                values.pop()
            if any(values):
                rows.append(values)
                blocks.append(TextBlock(" | ".join(values), {"sheet": sheet.title, "row": row_index}))
        if rows:
            tables.append(rows)
    return _finalize(
        path,
        "openpyxl",
        0.92,
        blocks,
        tables=tables,
        metadata={"sheet_names": workbook.sheetnames},
    )


_READERS = {
    ".txt": _read_txt,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".xlsx": _read_xlsx,
}


def read_attachment(path: str | Path) -> DocumentReadResult:
    """Read one attachment and return a typed success or failure result."""
    file_path = Path(path)
    if not file_path.exists() or not file_path.is_file():
        return _failure(file_path, "none", "UNREADABLE", "MISSING_FILE")
    if file_path.stat().st_size == 0:
        return _failure(file_path, "none", "EMPTY", "EMPTY_FILE")
    reader = _READERS.get(file_path.suffix.lower())
    if reader is None:
        return _failure(file_path, "none", "UNSUPPORTED", "UNSUPPORTED_FORMAT")
    try:
        return reader(file_path)
    except Exception as exc:  # parser errors become data, not crashes
        result = _failure(file_path, reader.__name__.removeprefix("_read_"), "UNREADABLE", "PARSER_ERROR")
        result.metadata["error_type"] = type(exc).__name__
        return result
