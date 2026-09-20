#!/usr/bin/env python3
"""Inspect representative document formats and candidate edge cases without answer labels."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path


def load(root: Path):
    return [json.loads(p.read_text(encoding="utf-8")) for p in sorted((root / "inbox").glob("email_*.json"))]


def preview_text(path: Path, limit=900):
    return re.sub(r"\s+", " ", path.read_text(encoding="utf-8", errors="replace"))[:limit]


def inspect_docx(path: Path):
    try:
        from docx import Document
        doc = Document(path)
        text = " ".join(p.text for p in doc.paragraphs)
        tables = [" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
        return {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "preview": re.sub(r"\s+", " ", (text + " " + " ".join(tables)))[:900]}
    except Exception as exc:
        return {"error": type(exc).__name__ + ": " + str(exc)}


def inspect_xlsx(path: Path):
    try:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        sheets = {}
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(max_row=12, values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    rows.append(" | ".join(values))
            sheets[sheet.title] = rows[:12]
        return {"sheets": list(book.sheetnames), "preview": sheets}
    except Exception as exc:
        return {"error": type(exc).__name__ + ": " + str(exc)}


def inspect_pdf(path: Path):
    try:
        import subprocess
        result = subprocess.run(["pdftotext", str(path), "-"], capture_output=True, text=True, timeout=20)
        return {"returncode": result.returncode, "preview": re.sub(r"\s+", " ", result.stdout)[:900], "stderr": result.stderr[:300]}
    except Exception as exc:
        return {"error": type(exc).__name__ + ": " + str(exc)}


def main():
    root = Path(sys.argv[1] if len(sys.argv) > 1 else "data/participant/sdoc-hackathon-bundle")
    records = load(root)
    by_id = {record["email_id"]: record for record in records}

    print("REPRESENTATIVE_EMAILS")
    for email_id in ["email_001", "email_004", "email_005", "email_055", "email_059", "email_070", "email_072"]:
        record = by_id.get(email_id)
        if record:
            print(json.dumps({"email_id": email_id, "subject": record.get("subject"), "body": re.sub(r"\s+", " ", record.get("body", ""))[:500], "attachments": record.get("attachments")}, ensure_ascii=False))

    print("FORMAT_SAMPLES")
    for filename in ["email_004_SI.txt", "email_004_BL.txt", "email_059_SI.pdf", "email_059_BL.pdf", "email_055_SI.xlsx", "email_055_BL.docx"]:
        path = root / "attachments" / filename
        if not path.exists():
            print(json.dumps({"file": filename, "missing": True}))
            continue
        ext = path.suffix.lower()
        if ext == ".txt":
            details = {"preview": preview_text(path)}
        elif ext == ".pdf":
            details = inspect_pdf(path)
        elif ext == ".docx":
            details = inspect_docx(path)
        elif ext == ".xlsx":
            details = inspect_xlsx(path)
        else:
            details = {"unsupported": True}
        print(json.dumps({"file": filename, "bytes": path.stat().st_size, **details}, ensure_ascii=False))

    candidates = []
    for record in records:
        attachments = record.get("attachments") or []
        body = record.get("body", "") or ""
        subject = record.get("subject", "") or ""
        reasons = []
        if not attachments:
            reasons.append("no_attachments")
        if any(Path(a).suffix.lower() in {".pdf", ".docx", ".xlsx"} for a in attachments):
            reasons.append("non_txt_format")
        combined = (subject + " " + body).lower()
        for marker in ["packing list", "commercial invoice", "certificate of origin", "tba", "missing", "unable to read", "blank"]:
            if marker in combined:
                reasons.append("text:" + marker.replace(" ", "_"))
        missing_paths = [a for a in attachments if not (root / a).exists()]
        if missing_paths:
            reasons.append("missing_path")
        if reasons:
            candidates.append({"email_id": record.get("email_id"), "subject": subject, "attachments": attachments, "reasons": reasons})
    print("EDGE_CASE_CANDIDATES count=" + str(len(candidates)))
    for item in candidates[:20]:
        print(json.dumps(item, ensure_ascii=False))


if __name__ == "__main__":
    main()
